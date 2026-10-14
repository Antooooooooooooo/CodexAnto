"""Parcours réels dans Chromium, sur un classeur synthétique (aucune donnée client)."""
import functools
import json
import tempfile
import threading
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

from openpyxl import Workbook
from playwright.sync_api import sync_playwright, expect

ROOT = Path(__file__).resolve().parents[1]

class QuietHandler(SimpleHTTPRequestHandler):
    def log_message(self, *_):
        pass

def fixture(path):
    wb = Workbook()
    ws = wb.active
    ws.title = 'DEGSE'
    ws.append(['PORTEFEUILLE FICTIF — VALIDATION TECHNIQUE'])
    ws.append(['SIREN', 'Raison Sociale', 'Rattachement commercial', 'Famille client', 'Nom et prénom VRC', 'Macro-sect. Nat. Libellé', 'CA GLOBAL 12 mois', 'Pays accompagnés par Orange'])
    ws.append(['037320595', 'Client historique fictif', 'AURA', 'MDM', 'Commercial A', 'Industrie', 12000, ''])
    ws.append(['123456789', 'Filiale fictive internationale', 'RM', 'SPES', '=FORMULE_FICTIVE', 'Services', 54000, 'Chine'])
    ws.cell(4, 5).data_type = 's'  # Texte potentiellement interprété comme formule lors d’un export CSV.
    ws.append(['987654321', '<img src=x onerror="window.injected=true">', 'AURA', 'HDM', 'Commercial B', 'Commerce', None, ''])
    ws.append(['123456789', 'DOUBLON FICTIF', 'AURA', 'MDM', 'Autre', 'Autre', 999999, ''])
    ws.append(['INVALIDE', 'Ligne invalide', 'RM', 'MDM'])
    for col, header in enumerate(['CUID','CUID','CUID Cible','CUID Cible','RDV BM.','RDV BM.'], start=9):
        ws.cell(2, col, header)
        ws.cell(4, col, f'VALEUR DISTINCTE {col}')
    details = wb.create_sheet('Détail Clients')
    details.append(['ENT. SIREN', 'Raison Sociale', 'GRP. Raison Sociale', 'GRP. SIREN Maison Mère'])
    details.append(['123456789', 'Filiale fictive internationale', 'Groupe mentionné dans Excel', '222222222'])
    details.cell(1, 5, 'CUID')
    details.cell(1, 6, 'CUID')
    details.cell(2, 5, 'DETAIL PREMIER')
    details.cell(2, 6, 'DETAIL SECOND')
    wb.save(path)

def main():
    server = ThreadingHTTPServer(('127.0.0.1', 0), functools.partial(QuietHandler, directory=str(ROOT)))
    threading.Thread(target=server.serve_forever, daemon=True).start()
    url = f'http://127.0.0.1:{server.server_port}/analyseurcode.html'
    checks = []
    try:
        with tempfile.TemporaryDirectory(prefix='degse-test-') as tmp, sync_playwright() as p:
            tmp = Path(tmp); xlsx = tmp / 'portefeuille-fictif.xlsx'; fixture(xlsx)
            browser = p.chromium.launch(executable_path='/usr/bin/chromium', args=['--no-sandbox'])
            context = browser.new_context(viewport={'width':1440,'height':1000}, accept_downloads=True)
            page = context.new_page(); errors, external = [], []
            page.on('pageerror', lambda err: errors.append(str(err)))
            page.on('request', lambda req: external.append(req.url) if req.url.startswith(('https://','http://')) and not req.url.startswith(url.rsplit('/',1)[0]) else None)
            page.goto(url)
            expect(page.locator('#metrics .metric strong').nth(0)).to_have_text('721')
            expect(page.locator('html')).to_have_attribute('data-theme', 'dark')
            expect(page.locator('#themeToggle')).to_have_attribute('aria-pressed', 'true')
            page.locator('#themeToggle').click()
            expect(page.locator('html')).to_have_attribute('data-theme', 'light')
            page.reload()
            expect(page.locator('html')).to_have_attribute('data-theme', 'light')
            page.locator('#themeToggle').click()
            expect(page.locator('html')).to_have_attribute('data-theme', 'dark')
            checks.append('Thème sombre par défaut, retour au clair et préférence conservée après rechargement')
            expect(page.locator('#metrics .metric strong').nth(1)).to_have_text('11')
            expect(page.locator('#targetCountries input:checked')).to_have_count(5)
            assert page.locator('#methodDialog, #countryFilters, #status, #analysis').count() == 0
            assert page.locator('.orange-logo').evaluate('(el) => el.complete && el.naturalWidth > 0')
            expect(page.locator('.brand-right')).to_contain_text('Analyseur interne')
            checks.append('Orange : logo inclus, ciblage unique, onze cibles documentées sans relecture forcée')
            page.locator('#view').select_option('confirmed')
            expect(page.locator('#tableBody tr')).to_have_count(11)
            expect(page.locator('#tableBody')).to_contain_text('KONE')
            page.locator('#search').fill('378317721')
            page.locator('[data-company="378317721"]').first.click()
            expect(page.locator('#summaryPanel')).to_be_visible()
            expect(page.locator('#editPanel')).not_to_be_visible()
            expect(page.locator('#researchOverview')).to_contain_text('Allemagne')
            page.locator('#editTab').click(); form = page.locator('#researchForm')
            expect(form.locator('[name="revenueAmount"]')).to_have_value('6405874000')
            expect(form.locator('[name="countryBreakdownKnown"]')).to_be_checked()
            expect(page.locator('#suppliedDetail')).to_contain_text('lot 16, ligne 474')
            expect(page.locator('#registryDetail')).to_contain_text('CA des comptes de la société')
            expect(page.locator('#suppliedDetail')).to_contain_text('Recherche documentaire')
            page.locator('[data-close="companyDialog"]').first.click()
            page.locator('#view').select_option('all');page.locator('#search').fill('418096392')
            page.locator('[data-company="418096392"]').first.click();page.locator('#editTab').click()
            expect(page.locator('#researchForm [name="revenueAmount"]')).to_have_value('')
            assert '449,4' in page.locator('#researchForm [name="franceRevenue"]').input_value()
            page.locator('[data-close="companyDialog"]').first.click()
            checks.append('Fiche en trois sections, CA mondial distinct du CA France et des comptes sociaux, lots et sources conservés')
            page.locator('#search').fill('');page.locator('#excelInput').set_input_files(str(xlsx))
            expect(page.locator('#fileMeta')).to_contain_text('3 entreprises importées')
            page.locator('#scope').select_option('portfolio');page.locator('#view').select_option('all')
            expect(page.locator('#tableBody tr')).to_have_count(3)
            expect(page.locator('#diagnosticContent')).to_contain_text('1 SIREN en doublon')
            expect(page.locator('#diagnosticContent')).to_contain_text('En-têtes répétés dans DEGSE : CUID, CUID Cible, RDV BM.')
            expect(page.locator('#diagnosticContent')).to_contain_text('En-têtes répétés dans Détail Clients')
            assert page.evaluate('window.injected') is None
            checks.append('Import XLSX : en-tête décalé, zéro initial, doublon, en-têtes répétés et affichage sécurisé')
            page.locator('#importDiagnostic').evaluate('(el) => el.open = true')
            page.locator('#mapDirection').select_option('Rattachement commercial')
            page.locator('#mapSegment').select_option('Famille client');page.locator('#applyMapping').click()
            page.locator('.advanced').evaluate('(el) => el.open = true')
            page.locator('#direction').select_option('RM');page.locator('#segment').select_option('SPES')
            expect(page.locator('#tableBody tr')).to_have_count(1)
            expect(page.locator('#tableBody')).to_contain_text('RM · SPES')
            page.locator('[data-company="123456789"]').first.click()
            expect(page.locator('#excelDetail')).to_contain_text('Groupe mentionné dans Excel')
            expect(page.locator('#researchOverview')).not_to_contain_text('Groupe mentionné dans Excel')
            expect(page.locator('#rawExcel')).to_contain_text('Détail Clients')
            for col in range(9,15):expect(page.locator('#rawExcel')).to_contain_text(f'VALEUR DISTINCTE {col}')
            expect(page.locator('#rawExcel')).to_contain_text('DETAIL PREMIER')
            expect(page.locator('#rawExcel')).to_contain_text('DETAIL SECOND')
            checks.append('Origines Excel séparées : AURA/RM, MDM/HDM/SPES, groupe et toutes les colonnes originales')
            page.locator('#editTab').click(); form = page.locator('#researchForm')
            values={'checkedAt':'2026-01-12','sector':'Services numériques FICTIFS','sectorSource':'https://example.com/sector','hq':'France','hqSource':'https://example.com/hq','groupName':'Groupe FICTIF','groupHQ':'Finlande','groupSource':'https://example.com/group','countries':'Italie ; Suisse ; Allemagne ; Chine ; Pays-Bas','presenceSource':'https://example.com/locations','revenueAmount':'200000000','revenueYear':'2024','revenueSource':'https://example.com/report','notes':'Données synthétiques pour test uniquement.'}
            for name,value in values.items():form.locator(f'[name="{name}"]').fill(value)
            for name,value in {'groupRelation':'Oui','presenceScope':'Groupe','presenceType':'Implantation','revenueCurrency':'EUR','revenueScope':'Groupe'}.items():form.locator(f'[name="{name}"]').select_option(value)
            form.locator('[name="reviewed"]').check();form.get_by_role('button',name='Enregistrer l’analyse').click()
            expect(page.locator('#companyDialog')).not_to_be_visible()
            expect(page.locator('#tableBody')).to_contain_text('Cible confirmée')
            expect(page.locator('#tableBody')).to_contain_text('Finlande')
            expect(page.locator('#tableBody')).to_contain_text('Relation Orange')
            checks.append('Saisie sourcée relue qualifiée sans modifier l’Excel ; siège entité et groupe distincts')
            page.locator('#groupRule').select_option('entity')
            expect(page.locator('#metrics .metric strong').nth(1)).to_have_text('0')
            page.locator('#groupRule').select_option('include')
            expect(page.locator('#metrics .metric strong').nth(1)).to_have_text('1')
            for checkbox in page.locator('#targetCountries input').all():
                if checkbox.is_checked() and checkbox.get_attribute('value') != 'Luxembourg':checkbox.uncheck()
            expect(page.locator('#tableBody')).not_to_contain_text('Cible confirmée')
            page.locator('#targetCountries input[value="International — Autre"]').check()
            page.locator('#targetCountries input[value="Luxembourg"]').uncheck()
            expect(page.locator('#tableBody')).to_contain_text('Cible confirmée')
            page.locator('#resetFilters').click();page.locator('#scope').select_option('portfolio');page.locator('#view').select_option('all')
            checks.append('Une seule sélection pays et critères en ligne : recalcul effectif, groupe et International Autre')
            page.locator('[data-select="123456789"]').check()
            expect(page.locator('#metrics .metric strong').nth(3)).to_have_text('1')
            page.locator('[data-company="123456789"][data-sales]').click()
            expect(page.locator('#salesPanel')).to_be_visible()
            commercial=page.locator('#commercialForm')
            commercial.locator('[name="commercialStatus"]').select_option('Contacté')
            commercial.locator('[name="nextActionDate"]').fill('2000-01-01')
            commercial.locator('[name="commercialNotes"]').fill('Relance fictive : valider la couverture en Suisse.')
            commercial.get_by_role('button',name='Enregistrer le suivi').click()
            expect(page.locator('#copyStatus')).to_have_text('Suivi enregistré.')
            brief=page.locator('#commercialBrief').input_value()
            assert 'Commercial / portefeuille' in brief and 'Chine' in brief and 'example.com/locations' in brief
            page.locator('#copyBriefBtn').click();expect(page.locator('#copyStatus')).to_contain_text('copi')
            page.locator('[data-close="companyDialog"]').first.click()
            page.locator('#view').select_option('due');expect(page.locator('#tableBody tr')).to_have_count(1)
            expect(page.locator('#tableBody')).to_contain_text('Relance à faire')
            page.locator('#view').select_option('selected');expect(page.locator('#tableBody tr')).to_have_count(1)
            checks.append('Sélection, suivi commercial, relances, préparation et copie de la synthèse')
            page.locator('#view').select_option('all');page.locator('#selectViewBtn').click()
            expect(page.locator('#metrics .metric strong').nth(3)).to_have_text('3')
            page.locator('#view').select_option('selected');expect(page.locator('#tableBody tr')).to_have_count(3)
            with page.expect_download() as d:page.locator('#exportBtn').click()
            d.value.save_as(str(tmp/'vue.csv'));csv=(tmp/'vue.csv').read_text(encoding='utf-8-sig')
            assert "'=FORMULE_FICTIVE" in csv and '200000000' in csv and '54000' in csv and 'Prospection - Prochaine action' in csv and '2000-01-01' in csv
            page.locator('.storage-menu').evaluate('(el) => el.open = true')
            with page.expect_download() as d:page.locator('#backupBtn').click()
            d.value.save_as(str(tmp/'backup.json'));backup=json.loads((tmp/'backup.json').read_text())
            assert backup['researches']['123456789']['reviewed'] is True
            assert backup['researches']['987654321']['selected'] is True and backup['researches']['987654321']['analyzed'] is False
            assert sum(sum(not e.get('kind') for e in r.get('contributions',[])) for r in backup['researches'].values()) == 16
            assert sum(r.get('registryState') == 'Fermée' for r in backup['researches'].values()) == 6
            assert sum(bool(r.get('registeredRevenue')) for r in backup['researches'].values()) == 588
            checks.append('Sélection de toute la vue, CSV professionnel protégé des formules, sauvegarde des comptes non analysés sélectionnés')
            page.reload();page.locator('#view').select_option('selected')
            expect(page.locator('#tableBody tr')).to_have_count(3)
            page.locator('#search').fill('123456789');expect(page.locator('#tableBody')).to_contain_text('Cible confirmée')
            expect(page.locator('#tableBody')).to_contain_text('Contacté')
            expect(page.locator('#tableBody')).to_contain_text('Filiale fictive internationale')
            page.locator('#researchInput').set_input_files(str(tmp/'backup.json'))
            expect(page.locator('#notice')).to_contain_text('Les saisies restent à relire')
            expect(page.locator('#tableBody')).to_contain_text('À qualifier')
            expect(page.locator('#storageStatus')).to_contain_text('Modifications enregistrées')
            stored=page.evaluate('localStorage.getItem("degse-international-v8")')
            assert json.loads(stored)['encoding'] == 'gzip-base64' and len(stored)<1500000
            page.reload();page.locator('#view').select_option('selected');page.locator('#search').fill('123456789')
            expect(page.locator('#tableBody')).to_contain_text('À qualifier')
            expect(page.locator('#tableBody')).to_contain_text('Contacté')
            checks.append('Persistance du suivi et de la sélection, compression sous le quota, restauration avec relecture des saisies')
            bad=tmp/'bad.json';bad.write_text(json.dumps({'schemaVersion':1,'researches':{'123456789':{'siren':'bad'}}}))
            page.locator('#researchInput').set_input_files(str(bad));expect(page.locator('#notice')).to_contain_text('Aucune recherche modifiée')
            page.locator('#search').fill('');page.locator('#excelInput').set_input_files(str(xlsx));expect(page.locator('#fileMeta')).to_contain_text('3 entreprises importées')
            bad_excel=Workbook();bad_excel.active.title='Autre';bad_excel.save(tmp/'bad.xlsx')
            page.locator('#excelInput').set_input_files(str(tmp/'bad.xlsx'));expect(page.locator('#notice')).to_contain_text('DEGSE est requis')
            page.locator('#scope').select_option('portfolio');page.locator('#view').select_option('all');expect(page.locator('#tableBody tr')).to_have_count(3)
            checks.append('Import incompatible et sauvegarde invalide rejetés sans perte du portefeuille ou des recherches')
            page.screenshot(path='/tmp/degse-orange-desktop-final.png',full_page=True)
            page.set_viewport_size({'width':390,'height':844});assert page.evaluate('document.documentElement.scrollWidth <= innerWidth')
            page.screenshot(path='/tmp/degse-orange-mobile-final.png',full_page=True)
            assert errors == [], errors; assert external == [], external
            checks.append('Présentation desktop/mobile, logo autonome, aucune erreur JavaScript ni requête externe')
            single=Workbook();ws=single.active;ws.title='DEGSE';ws.append(['SIREN','Raison Sociale','Direction commerciale','Marché']);ws.append(['444444444','NOUVEAU CLIENT FICTIF','RM','HDM']);single.save(tmp/'second.xlsx')
            page.locator('#excelInput').set_input_files(str(tmp/'second.xlsx'));expect(page.locator('#fileMeta')).to_contain_text('1 entreprises importées')
            page.locator('#scope').select_option('portfolio');page.locator('#view').select_option('all')
            page.locator('[data-company="444444444"]').first.click()
            expect(page.locator('#excelDetail dt').filter(has_text='Groupe mentionné dans Excel').locator('..').locator('dd')).to_have_text('Non renseigné')
            checks.append('Réimport : aucune information commerciale ou de groupe du classeur précédent réutilisée')
            legacy_context=browser.new_context();legacy_page=legacy_context.new_page();legacy_page.goto(url)
            prior_format=dict(backup['researches']['123456789'])
            for field in ['sourceDates','documentedCountries','operationalCountries','selected','commercialStatus','nextActionDate','commercialNotes']:prior_format.pop(field,None)
            legacy_page.evaluate('(payload) => localStorage.setItem("degse-international-v8", JSON.stringify(payload))',{'schemaVersion':1,'researches':{'123456789':prior_format}})
            legacy_page.reload();legacy_page.locator('#view').select_option('all');legacy_page.locator('#search').fill('123456789')
            expect(legacy_page.locator('#tableBody')).to_contain_text('Cible confirmée')
            legacy_page.locator('.advanced').evaluate('(el) => el.open = true');legacy_page.locator('#groupRule').select_option('entity')
            expect(legacy_page.locator('#storageStatus')).to_contain_text('Modifications enregistrées')
            assert json.loads(legacy_page.evaluate('localStorage.getItem("degse-international-v8")'))['encoding'] == 'gzip-base64'
            legacy_context.close();checks.append('Compatibilité des anciens JSON locaux et migration vers le nouveau suivi sans perte')
            browser.close()
        print(json.dumps({'passed':len(checks),'checks':checks},ensure_ascii=False,indent=2))
    finally:server.shutdown();server.server_close()

if __name__ == '__main__':main()
